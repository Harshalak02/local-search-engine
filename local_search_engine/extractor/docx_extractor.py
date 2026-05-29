# extractor/docx_extractor.py
# ─────────────────────────────────────────────────────────────
# Responsibility : extract text and metadata from DOCX files.
# Member A owns this file.
# ─────────────────────────────────────────────────────────────

import os
from pathlib import Path

from docx import Document

from .base_extractor import BaseExtractor

import sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from shared_types import ExtractedDocument


class DOCXExtractor(BaseExtractor):

    def can_handle(self, file_path: str) -> bool:
        return Path(file_path).suffix.lower() == ".docx"

    def extract(self, file_path: str) -> ExtractedDocument:
        doc = Document(file_path)
        metadata = {"file_type": "docx"}

        # core_properties holds author, title, created, modified, etc.
        props = doc.core_properties
        metadata["title"]         = props.title or Path(file_path).stem
        metadata["author"]        = props.author or ""
        metadata["modified_time"] = os.path.getmtime(file_path)

        # extract non-empty paragraphs
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        raw_text = "\n\n".join(paragraphs)

        chunks = self._chunk_text(raw_text)

        return ExtractedDocument(
            file_path=file_path,
            raw_text=raw_text,
            chunks=chunks,
            metadata=metadata,
        )
